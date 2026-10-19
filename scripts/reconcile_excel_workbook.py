import openpyxl
from pymongo import MongoClient
import json
import os

client = MongoClient('mongodb://127.0.0.1:27017')
auth_db = client['medimind_auth']
users_in_db = list(auth_db.users.find({}))
user_by_email = {u['email'].lower().strip(): u for u in users_in_db}

wb_path = r'Documents/MediMind_Role_Credentials.xlsx'
wb = openpyxl.load_workbook(wb_path, data_only=True)

print("=== EXCEL CREDENTIALS WORKBOOK AUDIT ===")
print(f"Workbook Path: {wb_path}")
print(f"Sheet Names: {wb.sheetnames}")

expected_sheets = [
    'Instructions',
    'Credential Index',
    'Family',
    'Doctor',
    'Department Head',
    'Hospital Admin',
    'Chairman'
]

sheet_check = {s: s in wb.sheetnames for s in expected_sheets}
print(f"Expected sheets present: {sheet_check} (All present: {all(sheet_check.values())})")

# Reconcile accounts from workbook
wb_accounts = []
duplicates_in_wb = []
seen_emails = set()

sheet_role_map = {
    'Family': 'FAMILY',
    'Doctor': 'DOCTOR',
    'Department Head': 'DEPARTMENT_HEAD',
    'Hospital Admin': 'HOSPITAL_ADMIN',
    'Chairman': 'CHAIRMAN'
}

sheet_counts = {}

for s in wb.sheetnames:
    if s in ['Instructions', 'Credential Index']:
        continue
    ws = wb[s]
    rows = list(ws.iter_rows(values_only=True))
    if not rows or not rows[0]:
        continue
    h = [str(x).strip().lower() for x in rows[0] if x is not None]
    e_i = next((i for i, x in enumerate(h) if 'email' in x or 'username' in x), None)
    id_i = next((i for i, x in enumerate(h) if 'account_id' in x or 'id' in x), None)
    role_i = next((i for i, x in enumerate(h) if 'role' in x), None)
    
    count = 0
    for r in rows[1:]:
        if not r or e_i is None or len(r) <= e_i or not r[e_i]:
            continue
        email = str(r[e_i]).strip().lower()
        if email in seen_emails:
            duplicates_in_wb.append(email)
        seen_emails.add(email)
        acc_id = str(r[id_i]).strip() if id_i is not None and len(r) > id_i and r[id_i] else ''
        role = str(r[role_i]).strip() if role_i is not None and len(r) > role_i and r[role_i] else sheet_role_map.get(s, s)
        wb_accounts.append({
            'sheet': s,
            'email': email,
            'account_id': acc_id,
            'role': role
        })
        count += 1
    sheet_counts[s] = count

print(f"\nAccount counts per sheet:")
for s, c in sheet_counts.items():
    print(f"  {s:20s}: {c} accounts")
print(f"Total credential accounts in workbook: {len(wb_accounts)}")
print(f"Duplicates in workbook: {duplicates_in_wb}")

# Cross-reconcile with database
print(f"\nTotal canonical users in database: {len(users_in_db)}")
matched_in_db = 0
missing_in_db = []
role_mismatches = []

for acc in wb_accounts:
    email = acc['email']
    if email in user_by_email:
        matched_in_db += 1
        db_u = user_by_email[email]
        db_role = db_u.get('role')
        if db_role != sheet_role_map.get(acc['sheet']):
            role_mismatches.append({'email': email, 'wb_role': acc['role'], 'db_role': db_role})
    else:
        missing_in_db.append(email)

extra_in_db = [e for e in user_by_email.keys() if e not in seen_emails]

print(f"Matched in database: {matched_in_db} / {len(wb_accounts)}")
print(f"Missing in database: {missing_in_db}")
print(f"Extra accounts in database: {extra_in_db}")
print(f"Role mismatches: {role_mismatches}")

# Search for separate staff inventory workbook in Documents
print("\n--- Secondary Staff Workbook Inspection in Documents/ ---")
doc_files = os.listdir('Documents')
staff_workbooks = [f for f in doc_files if f.endswith(('.xlsx', '.xls', '.csv')) and f != 'MediMind_Role_Credentials.xlsx']
print(f"Staff or additional workbooks in Documents/: {staff_workbooks}")

audit_summary = {
    'workbook_file': wb_path,
    'total_workbook_accounts': len(wb_accounts),
    'sheet_breakdown': sheet_counts,
    'duplicates_in_workbook': duplicates_in_wb,
    'db_users_count': len(users_in_db),
    'matched_count': matched_in_db,
    'missing_in_db': missing_in_db,
    'extra_in_db': extra_in_db,
    'role_mismatches': role_mismatches,
    'secondary_staff_workbook_found': len(staff_workbooks) > 0,
    'secondary_staff_workbooks': staff_workbooks
}

with open('scripts/excel_reconciliation_audit.json', 'w') as f:
    json.dump(audit_summary, f, indent=2)
print("\nSaved excel reconciliation audit to scripts/excel_reconciliation_audit.json")
