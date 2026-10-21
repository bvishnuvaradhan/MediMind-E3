import openpyxl
import requests
import pymongo

client = pymongo.MongoClient('mongodb://localhost:27017')
doc_db = client['medimind_doctor']
auth_db = client['medimind_auth']

# 1. Load operational 1..29 credentials from workbook
wb = openpyxl.load_workbook('Documents/MediMind_Role_Credentials.xlsx', read_only=True)
ws = wb['Doctor']
headers = [c for c in next(ws.iter_rows(values_only=True))]
email_idx = headers.index('Email / Username')
pw_idx = headers.index('Password')

doctor_creds = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[email_idx]:
        doctor_creds.append({
            'index': len(doctor_creds) + 1,
            'roster_id': f"DOC-{len(doctor_creds) + 1:03d}",
            'email': row[email_idx],
            'password': row[pw_idx],
            'name': row[headers.index('Full Name')],
            'hospital': row[headers.index('Hospital Name')],
            'department': row[headers.index('Department Name')]
        })

print(f"Loaded {len(doctor_creds)} operational doctor credentials from workbook.")

# 2. Add newly provisioned 30..66 doctors
import json
with open('scripts/historical_66_doctors.json', 'r', encoding='utf-8') as f:
    h66 = json.load(f)

existing_29_names = {c['name'] for c in doctor_creds}
candidates = [d for d in h66 if d.get('name') not in existing_29_names][:37]

for idx, d in enumerate(candidates, 30):
    name_parts = [p.lower() for p in d.get('name').replace('Dr.', '').strip().split() if p]
    clean_slug = '.'.join(name_parts)
    
    orig_hosp = d.get('hospitalName')
    if 'MediMind Central' in orig_hosp:
        hosp_slug = 'aarogyam'
        hosp_name = 'Aarogyam Multispeciality Hospital'
    elif 'Apex' in orig_hosp:
        hosp_slug = 'nexora'
        hosp_name = 'Nexora Advanced Care Hospital'
    else:
        hosp_slug = 'harborview'
        hosp_name = 'Harborview Medical Centre'
        
    email = f"dr.{clean_slug}@{hosp_slug}.hospital"
    pw = f"Doctor#2026${idx:02d}!Dev"
    
    doctor_creds.append({
        'index': idx,
        'roster_id': f"DOC-{idx:03d}",
        'email': email,
        'password': pw,
        'name': d.get('name'),
        'hospital': hosp_name,
        'department': d.get('departmentName')
    })

print(f"Total doctor credentials assembled: {len(doctor_creds)} (Expect 66)")

# Test all 66 logins directly on auth-service (port 5001)
successes = []
failures = []

for cred in doctor_creds:
    try:
        res = requests.post(
            'http://localhost:5001/api/auth/login',
            json={'email': cred['email'], 'password': cred['password']},
            timeout=5
        )
        if res.status_code == 200:
            data = res.json()
            user = data.get('data', {}).get('user', {})
            token = data.get('data', {}).get('token')
            if token and user.get('role') == 'DOCTOR':
                successes.append(cred)
            else:
                failures.append((cred, res.status_code, res.text))
        else:
            failures.append((cred, res.status_code, res.text))
    except Exception as e:
        failures.append((cred, 500, str(e)))

print("\n=== 66-Doctor Authentication Summary ===")
print(f"Total Accounts Tested: {len(doctor_creds)}")
print(f"Login Successes: {len(successes)} / {len(doctor_creds)}")
print(f"Login Failures: {len(failures)}")

if failures:
    print("\nFailures:")
    for f in failures:
        print(f"  {f[0]['roster_id']} | {f[0]['email']} | Status: {f[1]} | Err: {f[2][:100]}")
else:
    print("\nALL 66 DOCTOR ACCOUNTS ARE 100% LOGIN CAPABLE AND OPERATIONAL!")
