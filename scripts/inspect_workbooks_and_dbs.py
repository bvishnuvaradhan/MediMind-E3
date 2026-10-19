import openpyxl
import pymongo
import json
import os
import glob

print("=== 1. MONGODB DATABASE INVENTORY ===")
client = pymongo.MongoClient('mongodb://127.0.0.1:27017')
db_names = [
    'medimind_auth',
    'medimind_family',
    'medimind_doctor',
    'medimind_hospital',
    'medimind_appointment',
    'medimind_records',
    'medimind_knowledge',
    'medimind_ai'
]

db_summary = {}
for name in db_names:
    db = client[name]
    cols = db.list_collection_names()
    db_summary[name] = {}
    print(f"\nDatabase: {name}")
    for col in cols:
        cnt = db[col].count_documents({})
        db_summary[name][col] = cnt
        print(f"  - {col}: {cnt}")

print("\n=== 2. CREDENTIALS WORKBOOK (Documents/MediMind_Role_Credentials.xlsx) ===")
wb_path = 'Documents/MediMind_Role_Credentials.xlsx'
wb = openpyxl.load_workbook(wb_path, read_only=True)
print('Sheet names:', wb.sheetnames)
wb_summary = {}
for s in wb.sheetnames:
    ws = wb[s]
    rows = list(ws.iter_rows(values_only=True))
    non_empty = [r for r in rows if any(x is not None for x in r)]
    wb_summary[s] = len(non_empty)
    print(f"Sheet '{s}': {len(non_empty)} rows (Header + {len(non_empty)-1} data rows)")
    if non_empty:
        header = [str(x) if x is not None else '' for x in non_empty[0]]
        print(f"  Columns: {header}")

print("\n=== 3. SEARCHING FOR 66 DOCTOR PROFILES MENTION OR SEEDS ===")
# Search scripts and backend seed files for doctor definitions
py_js_files = glob.glob('backend/**/*.js', recursive=True) + glob.glob('scripts/**/*.py', recursive=True) + glob.glob('scripts/**/*.json', recursive=True) + glob.glob('Documents/**/*.txt', recursive=True)
matches_66 = []
for f in py_js_files:
    try:
        with open(f, 'r', encoding='utf-8', errors='ignore') as fh:
            content = fh.read()
            if '66' in content and 'doctor' in content.lower():
                matches_66.append(f)
    except Exception:
        pass
print("Files mentioning '66' and 'doctor':", matches_66)

print("\n=== 4. INSPECTING DOCTOR COLLECTION ===")
doc_db = client['medimind_doctor']
doctors = list(doc_db['doctors'].find())
print(f"Total doctors in medimind_doctor.doctors: {len(doctors)}")
hosp_counts = {}
dept_counts = {}
status_counts = {}
for d in doctors:
    h = str(d.get('hospital_id') or d.get('hospital_name') or 'Unknown')
    dept = str(d.get('department_id') or d.get('department_name') or 'Unknown')
    st = str(d.get('status', 'Unknown'))
    status_counts[st] = status_counts.get(st, 0) + 1
    hosp_counts[h] = hosp_counts.get(h, 0) + 1
    dept_counts[dept] = dept_counts.get(dept, 0) + 1

print("Doctor status breakdown:", status_counts)
print(f"Hospitals referenced ({len(hosp_counts)}):", hosp_counts)

print("\n=== 5. INSPECTING APPOINTMENT DISCREPANCY ===")
appt_db = client['medimind_appointment']
appts = list(appt_db['appointments'].find())
print(f"Total appointments in medimind_appointment.appointments: {len(appts)}")

appt_status = {}
appt_type = {}
appt_dates = []
for a in appts:
    st = a.get('status', 'Unknown')
    appt_status[st] = appt_status.get(st, 0) + 1
    typ = a.get('appointment_type', 'Unknown')
    appt_type[typ] = appt_type.get(typ, 0) + 1
    d = a.get('appointment_date')
    if d:
        appt_dates.append(str(d)[:10])

print("Appointment status breakdown:", appt_status)
print("Appointment type breakdown:", appt_type)
if appt_dates:
    appt_dates.sort()
    print(f"Appointment date range: {appt_dates[0]} to {appt_dates[-1]}")
    # Count appointments before vs after Aug 1, 2026
    before_aug = sum(1 for d in appt_dates if d < '2026-08-01')
    after_aug = sum(1 for d in appt_dates if d >= '2026-08-01')
    print(f"Appointments before 2026-08-01: {before_aug}")
    print(f"Appointments on/after 2026-08-01: {after_aug}")

print("\n=== 6. INSPECTING AI PREDICTIONS PER FAMILY MEMBER ===")
ai_db = client['medimind_ai']
fam_db = client['medimind_family']
members = list(fam_db['familymembers'].find())
predictions = list(ai_db['predictions'].find())
print(f"Total predictions in medimind_ai.predictions: {len(predictions)}")
print(f"Total family members in medimind_family.familymembers: {len(members)}")

member_pred_counts = {}
for m in members:
    mid = str(m['_id'])
    m_name = m.get('full_name')
    # Check predictions linked by family_member_id or patient_id
    linked = [p for p in predictions if str(p.get('family_member_id') or p.get('patient_id') or p.get('familyMemberId') or '') == mid]
    member_pred_counts[m_name] = len(linked)

print("Sample predictions per member (first 10):")
for k in list(member_pred_counts.keys())[:10]:
    print(f"  - {k}: {member_pred_counts[k]}")

unlinked_preds = [p for p in predictions if not any(str(p.get('family_member_id') or p.get('patient_id') or p.get('familyMemberId') or '') == str(m['_id']) for m in members)]
print(f"Predictions unlinked to any of the 44 family members: {len(unlinked_preds)}")

print("\n=== 7. INSPECTING HOSPITALS AND DEPARTMENTS ===")
hosp_db = client['medimind_hospital']
hospitals = list(hosp_db['hospitals'].find())
departments = list(hosp_db['departments'].find())
dept_heads = list(hosp_db['departmentheads'].find())
print(f"Hospitals: {len(hospitals)}")
for h in hospitals:
    print(f"  - {h.get('name')} (ID: {h.get('_id')}, code: {h.get('hospital_code')}, status: {h.get('status')})")
print(f"Departments: {len(departments)}")
print(f"Department Heads: {len(dept_heads)}")
