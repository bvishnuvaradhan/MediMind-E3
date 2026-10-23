import requests
import io
import time
import openpyxl

chair_login = requests.post(
    'http://localhost:5001/api/auth/login',
    json={'email': 'chairman@medimind.org', 'password': 'Chair#2026$01!Dev'}
)
token = chair_login.json()['data']['token']
headers = {'Authorization': f'Bearer {token}'}

report_types = ['doctor_appointment', 'hospital', 'user', 'appointment', 'ai']

print("=== TESTING EXCEL EXPORT & PREVIEW FOR ALL 5 REPORT TYPES ===")

for r_type in report_types:
    print(f"\n--- Testing Report Type: {r_type} ---")
    
    # 1. Preview
    t0 = time.time()
    res_prev = requests.get(f'http://localhost:5000/api/appointments/reports/preview?reportType={r_type}&hospitalId=ALL', headers=headers)
    t1 = time.time()
    assert res_prev.status_code == 200, f"Preview failed for {r_type}: {res_prev.status_code} {res_prev.text}"
    prev_data = res_prev.json().get('data', {})
    print(f"  [PASS] Preview 200 OK ({t1 - t0:.2f}s) - Title: {prev_data.get('title')}")
    print(f"         Summary: {prev_data.get('summary')}")
    
    # 2. Excel Export
    t0 = time.time()
    res_excel = requests.get(f'http://localhost:5000/api/appointments/reports/export-excel?reportType={r_type}&hospitalId=ALL', headers=headers)
    t1 = time.time()
    assert res_excel.status_code == 200, f"Export failed for {r_type}: {res_excel.status_code} {res_excel.text}"
    content_len = len(res_excel.content)
    print(f"  [PASS] Excel Export 200 OK ({t1 - t0:.2f}s) - Payload: {content_len} bytes")
    
    # 3. Workbook Validation with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(res_excel.content), data_only=True)
    sheets = wb.sheetnames
    print(f"         Sheets ({len(sheets)}): {sheets}")
    for s_name in sheets:
        sheet = wb[s_name]
        print(f"           - '{s_name}': {sheet.max_row} rows, {sheet.max_column} cols")

print("\n[ALL 5 REPORT TYPES VERIFIED 100% WORKING]")
