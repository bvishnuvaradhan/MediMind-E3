import requests
import time
import json
import openpyxl

wb = openpyxl.load_workbook('Documents/MediMind_Role_Credentials.xlsx', data_only=True)
creds = {}
for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows: continue
    h = [str(x).lower() if x else '' for x in rows[0]]
    e_idx, p_idx = None, None
    for i, col in enumerate(h):
        if 'email' in col or 'username' in col: e_idx = i
        if 'password' in col: p_idx = i
    if e_idx is not None and p_idx is not None:
        for r in rows[1:]:
            if r[e_idx] and r[p_idx]:
                creds[str(r[e_idx]).strip().lower()] = str(r[p_idx]).strip()

c_email = 'chairman@medimind.org'
res = requests.post('http://localhost:5000/api/auth/login', json={
    'email': c_email,
    'password': creds[c_email]
})
token = res.json()['data']['token']
headers = {'Authorization': f'Bearer {token}'}

print("=== MediMind Chairman Dashboard End-to-End Performance Benchmark ===")

# Benchmark 1: GET /api/appointments/stats
stats_latencies = []
for _ in range(5):
    t0 = time.perf_counter()
    r = requests.get('http://localhost:5000/api/appointments/stats', headers=headers)
    stats_latencies.append((time.perf_counter() - t0) * 1000)

print(f"1. Aggregated Stats API (/api/appointments/stats):")
print(f"   Payload: {len(r.content)} bytes ({len(r.content)/1024:.2f} KB)")
print(f"   Status: {r.status_code}")
print(f"   Latency: Min: {min(stats_latencies):.2f} ms | Avg: {sum(stats_latencies)/len(stats_latencies):.2f} ms | Max: {max(stats_latencies):.2f} ms")

# Benchmark 2: Concurrent Dashboard Data Fetch (Simulating frontend Promise.all)
# Endpoints: /api/appointments/stats, /api/hospitals, /api/departments, /api/doctors, /api/hospital-requests?status=Pending, /api/ai/analytics
import concurrent.futures

def fetch(url):
    t0 = time.perf_counter()
    resp = requests.get(url, headers=headers)
    dur = (time.perf_counter() - t0) * 1000
    return url, resp.status_code, len(resp.content), dur

urls = [
    'http://localhost:5000/api/appointments/stats',
    'http://localhost:5000/api/hospitals',
    'http://localhost:5000/api/departments',
    'http://localhost:5000/api/doctors',
    'http://localhost:5000/api/hospital-requests?status=Pending',
    'http://localhost:5000/api/ai/analytics',
]

dashboard_latencies = []
for _ in range(5):
    t_start = time.perf_counter()
    with concurrent.futures.ThreadPoolExecutor(max_workers=6) as executor:
        results = list(executor.map(fetch, urls))
    dashboard_latencies.append((time.perf_counter() - t_start) * 1000)

print(f"\n2. Full Concurrent Dashboard Load (Promise.all Simulation):")
print(f"   Total Endpoints: {len(urls)}")
for url, status, size, dur in results:
    path = url.split(':5000')[1]
    print(f"   - {path:45}: {dur:6.2f} ms | {size/1024:6.2f} KB | Status {status}")
print(f"   Overall Wall-Clock Time: Min: {min(dashboard_latencies):.2f} ms | Avg: {sum(dashboard_latencies)/len(dashboard_latencies):.2f} ms | Max: {max(dashboard_latencies):.2f} ms")

# Benchmark 3: Comparison Summary
unoptimized_baseline = 10831.68
optimized_avg = sum(dashboard_latencies)/len(dashboard_latencies)
speedup = unoptimized_baseline / optimized_avg
payload_before = 3620000
payload_after = sum(r[2] for r in results)

print(f"\n3. Optimization Summary:")
print(f"   - Baseline Wall-Clock Time (Sequential + Raw Appointments): {unoptimized_baseline:.2f} ms (~10.8 s)")
print(f"   - Optimized Wall-Clock Time (Concurrent + Aggregated Stats): {optimized_avg:.2f} ms")
print(f"   - Speedup Factor: {speedup:.1f}x faster")
print(f"   - Total Payload Before: {payload_before/1024:.2f} KB (~3.6 MB)")
print(f"   - Total Payload After:  {payload_after/1024:.2f} KB")
print(f"   - Payload Reduction:    {(1 - payload_after/payload_before)*100:.2f}%")
