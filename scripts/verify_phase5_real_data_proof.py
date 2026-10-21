import requests
import pymongo
import openpyxl
import io
import datetime

client = pymongo.MongoClient('mongodb://localhost:27017')
appt_db = client['medimind_appointment']
doc_db = client['medimind_doctor']
auth_db = client['medimind_auth']
hosp_db = client['medimind_hospital']

print("=====================================================================")
print("MediMind — Real-Data Proof, Anti-Mock Verification & Security Testing")
print("=====================================================================")

# A. Authenticate test personas
# 1. Chairman
chair_res = requests.post(
    'http://localhost:5001/api/auth/login',
    json={'email': 'chairman@medimind.org', 'password': 'Chair#2026$01!Dev'}
)
assert chair_res.status_code == 200, f"Chairman login failed: {chair_res.text}"
chair_token = chair_res.json()['data']['token']
chair_headers = {'Authorization': f'Bearer {chair_token}'}

# 2. Doctor (Unauthorized for reports)
doc_res = requests.post(
    'http://localhost:5001/api/auth/login',
    json={'email': 'dr.rahul.sharma@aarogyam.hospital', 'password': 'Doctor#2026$01!Dev'}
)
assert doc_res.status_code == 200, f"Doctor login failed: {doc_res.text}"
doc_token = doc_res.json()['data']['token']
doc_headers = {'Authorization': f'Bearer {doc_token}'}

# 3. Hospital Admin (Aarogyam)
ha_res = requests.post(
    'http://localhost:5001/api/auth/login',
    json={'email': 'admin@aarogyam.hospital', 'password': 'Admin#2026$01!Dev'}
)
assert ha_res.status_code == 200, f"Hospital Admin login failed: {ha_res.text}"
ha_token = ha_res.json()['data']['token']
ha_headers = {'Authorization': f'Bearer {ha_token}'}

# Test 1: Compare preview doctor counts against DB ground truth
print("\n[TEST 1] Doctor Counts: Preview vs Independent DB Ground Truth")
db_active_docs = doc_db.doctors.count_documents({'status': 'ACTIVE'})
db_total_docs = doc_db.doctors.count_documents({})
prev_res = requests.get('http://localhost:5000/api/appointments/reports/preview', headers=chair_headers)
assert prev_res.status_code == 200
prev_data = prev_res.json()['data']
print(f"  DB Active Doctors: {db_active_docs} | Preview Active Doctors: {prev_data['summary']['activeDoctors']}")
print(f"  DB Total Doctors:  {db_total_docs} | Preview Total Doctors:  {prev_data['summary']['totalDoctors']}")
assert db_active_docs == prev_data['summary']['activeDoctors'] == 66
assert db_total_docs == prev_data['summary']['totalDoctors'] == 68
print("  --> PASS: Doctor counts match database ground truth 100%.")

# Test 2: Compare appointment totals and status breakdown against DB
print("\n[TEST 2] Appointment Totals & Status Breakdown vs Independent DB")
db_total_appts = appt_db.appointments.count_documents({})
db_status_counts = {}
for s in ['COMPLETED', 'CONFIRMED', 'BOOKED', 'CANCELLED', 'RESCHEDULED']:
    db_status_counts[s] = appt_db.appointments.count_documents({'status': s})

prev_appts = prev_data['summary']['totalAppointments']
prev_breakdown = prev_data['summary']['statusBreakdown']
print(f"  DB Total: {db_total_appts} | Preview Total: {prev_appts}")
assert db_total_appts == prev_appts == 2851
for s in ['COMPLETED', 'CONFIRMED', 'BOOKED', 'CANCELLED', 'RESCHEDULED']:
    print(f"  Status {s:11s} -> DB: {db_status_counts[s]} | Preview: {prev_breakdown[s]}")
    assert db_status_counts[s] == prev_breakdown[s]
print("  --> PASS: Appointment totals and breakdown match database 100%.")

# Test 3: Download Excel via Gateway and verify all sheets & row counts
print("\n[TEST 3] Download Excel via Gateway and Verify All Sheets")
excel_res = requests.get('http://localhost:5000/api/appointments/reports/export-excel', headers=chair_headers)
assert excel_res.status_code == 200
assert excel_res.headers.get('Content-Type') == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
assert 'MediMind_Doctor_Appointment_Report_' in excel_res.headers.get('Content-Disposition')
wb = openpyxl.load_workbook(io.BytesIO(excel_res.content))
expected_sheets = ['Report Summary', 'Doctor Directory', 'Appointments', 'Doctor Appointment Summary']
assert wb.sheetnames == expected_sheets, f"Sheet mismatch: {wb.sheetnames}"
print(f"  Verified 4/4 sheets present: {wb.sheetnames}")
s2 = wb['Doctor Directory']
s3 = wb['Appointments']
s4 = wb['Doctor Appointment Summary']
print(f"  Doctor Directory rows: {s2.max_row} (1 header + 68 doctors in scope)")
print(f"  Appointments rows:     {s3.max_row} (1 header + 2,851 live appointments)")
print(f"  Doctor Summary rows:   {s4.max_row} (1 header + 68 doctors)")
assert s2.max_row == 69
assert s3.max_row == 2852
assert s4.max_row == 69
print("  --> PASS: Excel workbook downloaded and contains complete live dataset.")

# Test 4: Filter parity between Preview and Excel export (Facility Filter)
print("\n[TEST 4] Filter Parity Test: Aarogyam Multispeciality Hospital Scope")
aarogyam_id = str(hosp_db.hospitals.find_one({'code': 'MM-HYD-01'})['_id'])
db_aarogyam_appts = appt_db.appointments.count_documents({'hospital_id': hosp_db.hospitals.find_one({'code': 'MM-HYD-01'})['_id']})
db_aarogyam_docs = doc_db.doctors.count_documents({'hospital_id': hosp_db.hospitals.find_one({'code': 'MM-HYD-01'})['_id']})

# Preview with filter
f_prev = requests.get(f'http://localhost:5000/api/appointments/reports/preview?hospitalId={aarogyam_id}', headers=chair_headers).json()['data']
# Excel with filter
f_excel = requests.get(f'http://localhost:5000/api/appointments/reports/export-excel?hospitalId={aarogyam_id}', headers=chair_headers)
wb_f = openpyxl.load_workbook(io.BytesIO(f_excel.content))
f_s3 = wb_f['Appointments']
f_s2 = wb_f['Doctor Directory']

print(f"  Aarogyam DB Appts: {db_aarogyam_appts} | Preview: {f_prev['summary']['totalAppointments']} | Excel: {f_s3.max_row - 1}")
print(f"  Aarogyam DB Docs:  {db_aarogyam_docs}  | Preview: {f_prev['summary']['totalDoctors']}  | Excel: {f_s2.max_row - 1}")
assert db_aarogyam_appts == f_prev['summary']['totalAppointments'] == (f_s3.max_row - 1) == 597
assert db_aarogyam_docs == f_prev['summary']['totalDoctors'] == (f_s2.max_row - 1)
print("  --> PASS: Filter parity holds exactly between Preview, Excel Export, and DB (597 appointments).")

# Test 5: Status Filter Parity (COMPLETED appointments only)
print("\n[TEST 5] Status Filter Parity Test (COMPLETED only)")
c_prev = requests.get('http://localhost:5000/api/appointments/reports/preview?status=COMPLETED', headers=chair_headers).json()['data']
c_excel = requests.get('http://localhost:5000/api/appointments/reports/export-excel?status=COMPLETED', headers=chair_headers)
wb_c = openpyxl.load_workbook(io.BytesIO(c_excel.content))
c_s3 = wb_c['Appointments']
print(f"  Completed DB: {db_status_counts['COMPLETED']} | Preview: {c_prev['summary']['totalAppointments']} | Excel: {c_s3.max_row - 1}")
assert db_status_counts['COMPLETED'] == c_prev['summary']['totalAppointments'] == (c_s3.max_row - 1) == 1763
print("  --> PASS: Status filter correctly isolates exactly 1,763 completed appointments.")

# Test 6: Invalid Date Range Validation (startDate > endDate)
print("\n[TEST 6] Input Validation: Invalid Date Range (startDate > endDate)")
inv_res = requests.get('http://localhost:5000/api/appointments/reports/preview?startDate=2026-11-01&endDate=2026-08-01', headers=chair_headers)
print(f"  Status: {inv_res.status_code} | Msg: {inv_res.json().get('message')}")
assert inv_res.status_code == 400
assert 'startDate must be earlier than or equal to endDate' in inv_res.json().get('message')
print("  --> PASS: Invalid date range properly rejected with HTTP 400 Bad Request.")

# Test 7: Unauthorized Access Protection (Doctor Role Prohibited)
print("\n[TEST 7] Security / Authorization: Doctor Role Access Rejection")
d_prev = requests.get('http://localhost:5000/api/appointments/reports/preview', headers=doc_headers)
d_excel = requests.get('http://localhost:5000/api/appointments/reports/export-excel', headers=doc_headers)
print(f"  Doctor Preview Call Status: {d_prev.status_code} (Expect 403)")
print(f"  Doctor Export Call Status:  {d_excel.status_code} (Expect 403)")
assert d_prev.status_code == 403
assert d_excel.status_code == 403
print("  --> PASS: Doctor role strictly forbidden from executive report endpoints.")

# Test 8: Tenant Isolation / Hospital Admin Tamper Protection
print("\n[TEST 8] Tenant Isolation: Hospital Admin Tamper Protection")
harborview_id = str(hosp_db.hospitals.find_one({'code': 'MM-MUM-03'})['_id'])
# Aarogyam Admin attempts to query Harborview's data by tampering query param
ha_tamper = requests.get(f'http://localhost:5000/api/appointments/reports/preview?hospitalId={harborview_id}', headers=ha_headers)
assert ha_tamper.status_code == 200
ha_tamper_data = ha_tamper.json()['data']
# Must return ONLY Aarogyam's data (597 appointments), ignoring the tampered hospitalId!
print(f"  Aarogyam Admin Total Appointments returned: {ha_tamper_data['summary']['totalAppointments']} (Must be 597, NOT Harborview's 1,470)")
assert ha_tamper_data['summary']['totalAppointments'] == 597
print("  --> PASS: Tenant isolation strictly enforced; hospital admin cannot tamper query to access other hospitals.")

# Test 9: Spreadsheet Formula Injection Sanitization (CWE-1236)
print("\n[TEST 9] Spreadsheet Formula Injection Protection (CWE-1236)")
# Verify that all cells beginning with =, +, -, @ are escaped with leading single quote
formula_threats = 0
for sheet in wb.sheetnames:
    ws = wb[sheet]
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if isinstance(cell, str) and len(cell) > 1 and cell[0] in ['=', '+', '-', '@']:
                # If unescaped formula
                formula_threats += 1
print(f"  Formula injection threats found in exported workbook: {formula_threats}")
assert formula_threats == 0
print("  --> PASS: Formula sanitization verified. Zero unescaped spreadsheet execution vectors.")

# Test 10: Empty Result Set Handling
print("\n[TEST 10] Empty Result Set Handling")
empty_res = requests.get('http://localhost:5000/api/appointments/reports/preview?search=NonExistentSpecialistXYZ123', headers=chair_headers)
assert empty_res.status_code == 200
empty_data = empty_res.json()['data']
print(f"  Empty search returned: {empty_data['summary']['totalDoctors']} doctors, {empty_data['summary']['totalAppointments']} appointments.")
assert empty_data['summary']['totalDoctors'] == 0
assert len(empty_data['doctors']) == 0
print("  --> PASS: Empty result set handled cleanly without error or mock substitution.")

print("\n=====================================================================")
print("ALL 10 VERIFICATION GATES PASSED! REAL-DATA PROOF CONFIRMED 100%.")
print("=====================================================================")
