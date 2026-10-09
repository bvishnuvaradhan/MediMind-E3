import openpyxl
import bcrypt
from pymongo import MongoClient
from bson import ObjectId
import os
from dotenv import load_dotenv

load_dotenv('backend/.env')
c = MongoClient('mongodb://127.0.0.1:27017')
db_auth = c['medimind_auth']
db_hosp = c['medimind_hospital']
db_doc = c['medimind_doctor']
db_fam = c['medimind_family']
db_appt = c['medimind_appointment']
db_rec = c['medimind_records']
db_ai = c['medimind_ai']

# 1. Verify Excel Credentials against Database
wb = openpyxl.load_workbook(r'D:\MediMind_Backups\MediMind_Credentials_20261009.xlsx', data_only=True)
sheets = ['Chairman', 'Hospital Admin', 'Department Head', 'Doctor', 'Family']
print('Sheets in Excel:', wb.sheetnames)
assert set(sheets).issubset(set(wb.sheetnames)), 'Missing sheets in credentials workbook'

total_creds = 0
for s in sheets:
    ws = wb[s]
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]
    email_idx = headers.index('Email / Username')
    pwd_idx = headers.index('Password')
    role_idx = headers.index('Role')
    for row in rows[1:]:
        email = row[email_idx]
        pwd = row[pwd_idx]
        role = row[role_idx]
        user = db_auth.users.find_one({'email': email})
        assert user is not None, f'User {email} not found in DB!'
        assert user['role'] == role, f'Role mismatch for {email}'
        assert bcrypt.checkpw(pwd.encode('utf-8'), user['password_hash'].encode('utf-8')), f'Password check failed for {email}'
        total_creds += 1

print(f'Successfully validated {total_creds}/59 accounts against database password hashes!')

# 2. Check Referential Integrity
hosp_ids = {h['_id']: h for h in db_hosp.hospitals.find({})}
dept_ids = {d['_id']: d for d in db_hosp.departments.find({})}
doc_ids = {d['_id']: d for d in db_doc.doctors.find({})}
fam_ids = {f['_id']: f for f in db_fam.families.find({})}
mem_ids = {m['_id']: m for m in db_fam.familymembers.find({})}

# Pending hospitals isolation
for h_id, h in hosp_ids.items():
    if h.get('status') == 'PENDING':
        dept_cnt = db_hosp.departments.count_documents({'hospital_id': h_id})
        doc_cnt = db_doc.doctors.count_documents({'hospital_id': h_id})
        head_cnt = db_hosp.departmentheads.count_documents({'hospital_id': h_id})
        assert dept_cnt == 0, f"Pending hospital {h['name']} has {dept_cnt} departments!"
        assert doc_cnt == 0, f"Pending hospital {h['name']} has {doc_cnt} doctors!"
        assert head_cnt == 0, f"Pending hospital {h['name']} has {head_cnt} heads!"
print('Pending hospitals have zero operational departments, doctors, or heads.')

# Check doctors
for d_id, d in doc_ids.items():
    assert d['hospital_id'] in hosp_ids, f'Doctor {d_id} has invalid hospital_id'
    assert d['department_id'] in dept_ids, f'Doctor {d_id} has invalid department_id'
    assert hosp_ids[d['hospital_id']]['status'] == 'ACTIVE', f'Doctor in non-active hospital'

# Check appointments
for a in db_appt.appointments.find({}):
    assert a['doctor_id'] in doc_ids, 'Invalid doctor in appointment'
    assert a['hospital_id'] in hosp_ids, 'Invalid hospital in appointment'
    assert a['department_id'] in dept_ids, 'Invalid dept in appointment'
    assert a['family_member_id'] in mem_ids, 'Invalid member in appointment'
print('All appointments have 100% valid foreign references.')

# Check AI predictions
pred_count = 0
for p in db_ai.predictions.find({}):
    assert ObjectId(p['family_member_id']) in mem_ids, 'Prediction has invalid family_member_id'
    pred_count += 1
print(f'All {pred_count} AI predictions have 100% valid references.')
print('Targeted verification passed successfully!')
