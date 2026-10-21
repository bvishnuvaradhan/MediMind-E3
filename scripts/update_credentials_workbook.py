import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import json

wb = openpyxl.load_workbook('Documents/MediMind_Role_Credentials.xlsx')

# Styles
header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
demo_header_fill = PatternFill(start_color='991B1B', end_color='991B1B', fill_type='solid')
gray_fill = PatternFill(start_color='F3F4F6', end_color='F3F4F6', fill_type='solid')
cell_font = Font(name='Segoe UI', size=10)
bold_font = Font(name='Segoe UI', size=10, bold=True)
inactive_font = Font(name='Segoe UI', size=10, italic=True, color='6B7280')

thin_border = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    top=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB')
)

# 1. Create or replace 'Demonstration Accounts' sheet
if 'Demonstration Accounts' in wb.sheetnames:
    del wb['Demonstration Accounts']

ws_demo = wb.create_sheet(title='Demonstration Accounts')
ws_demo.views.sheetView[0].showGridLines = True

demo_headers = [
    'Account ID', 'Full Name', 'Email / Username', 'Role Code', 'Assigned Facility', 
    'Department / Scope', 'Password', 'Account Status', 'Login Authorization', 
    'Verification Date', 'Security Classification'
]

ws_demo.append(demo_headers)
for col_num in range(1, len(demo_headers) + 1):
    c = ws_demo.cell(row=1, column=col_num)
    c.font = header_font
    c.fill = demo_header_fill
    c.alignment = Alignment(horizontal='center', vertical='center')

demo_rows = [
    (
        '6ad8a8d0939031fc76bf6f7c', 'Demo Dr. Orthopedics (Inactive)', 'demo.doctor.ortho@medimind.test',
        'DOCTOR', 'Aarogyam Multispeciality Hospital', 'Orthopedics & Trauma',
        'NONE - LOGIN DISABLED', 'INACTIVE', 'HTTP 403 (BLOCKED)', '2026-10-21', 'Dev/Test Inactive Demo'
    ),
    (
        '6ad8a8d0939031fc76bf6f7e', 'Demo Dr. Diabetology (Inactive)', 'demo.doctor.diab@medimind.test',
        'DOCTOR', 'Nexora Advanced Care Hospital', 'Diabetology',
        'NONE - LOGIN DISABLED', 'INACTIVE', 'HTTP 403 (BLOCKED)', '2026-10-21', 'Dev/Test Inactive Demo'
    ),
    (
        '6ad8a8d0939031fc76bf6f80', 'Demo DeptHead Ortho (Inactive)', 'demo.depthead.ortho@medimind.test',
        'DEPARTMENT_HEAD', 'Aarogyam Multispeciality Hospital', 'Orthopedics & Trauma',
        'NONE - LOGIN DISABLED', 'INACTIVE', 'HTTP 403 (BLOCKED)', '2026-10-21', 'Dev/Test Inactive Demo'
    ),
    (
        '6ad8a8d0939031fc76bf6f82', 'Demo Admin Harborview (Inactive)', 'demo.admin.harborview@medimind.test',
        'HOSPITAL_ADMIN', 'Harborview Medical Centre', 'Hospital Administration',
        'NONE - LOGIN DISABLED', 'INACTIVE', 'HTTP 403 (BLOCKED)', '2026-10-21', 'Dev/Test Inactive Demo'
    ),
]

for row_data in demo_rows:
    ws_demo.append(row_data)

for r in range(2, len(demo_rows) + 2):
    for c in range(1, len(demo_headers) + 1):
        cell = ws_demo.cell(row=r, column=c)
        cell.font = inactive_font
        cell.border = thin_border
        if c in [7, 8, 9]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

# Column width auto-adjust
for col in ws_demo.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws_demo.column_dimensions[col_letter].width = max(max_len + 3, 14)

# 2. Add 'Historical 66 Capacity' sheet
if 'Historical 66 Capacity' in wb.sheetnames:
    del wb['Historical 66 Capacity']

ws_cap = wb.create_sheet(title='Historical 66 Capacity')
ws_cap.views.sheetView[0].showGridLines = True

cap_headers = [
    'Roster #', 'Full Name', 'Assigned Hospital', 'Specialization', 
    'Provisioning Status', 'Account Type', 'Login Credential'
]
ws_cap.append(cap_headers)
for col_num in range(1, len(cap_headers) + 1):
    c = ws_cap.cell(row=1, column=col_num)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center', vertical='center')

with open('scripts/historical_66_doctors.json', 'r', encoding='utf-8') as f:
    hist_docs = json.load(f)

for idx, doc in enumerate(hist_docs, 1):
    status = "ACTIVE_OPERATIONAL_ACCOUNT" if doc.get('is_provisioned') else "UNPROVISIONED_STAFF_CAPACITY"
    login = "Active in Doctor Sheet" if doc.get('is_provisioned') else "NONE - NO LOGIN PROFILE"
    ws_cap.append([
        f"DOC-ROSTER-{idx:03d}",
        doc.get('full_name') or doc.get('name'),
        doc.get('hospital_name') or doc.get('hospital'),
        doc.get('specialization') or 'General Medicine',
        status,
        "Operational" if doc.get('is_provisioned') else "Design Planning Only",
        login
    ])

for r in range(2, len(hist_docs) + 2):
    for c in range(1, len(cap_headers) + 1):
        cell = ws_cap.cell(row=r, column=c)
        cell.font = cell_font
        cell.border = thin_border
        if c in [1, 5, 6, 7]:
            cell.alignment = Alignment(horizontal='center', vertical='center')

for col in ws_cap.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws_cap.column_dimensions[col_letter].width = max(max_len + 3, 15)

# 3. Update 'Credential Index'
ws_idx = wb['Credential Index']
# Check if demo accounts row already present
existing_idx = [row[0] for row in ws_idx.iter_rows(values_only=True)]
if 'Demonstration Accounts' not in existing_idx:
    ws_idx.append([
        'Demonstration Accounts', 'DEMO_INACTIVE', 4,
        'Aarogyam, Nexora, Harborview (1 Dept Head, 2 Doctors, 1 Admin)',
        'Local Development & Testing', '0% (Login Prohibited - HTTP 403 Enforced)',
        '2026-10-21', 'Dev/Test Inactive Demo'
    ])
    r = ws_idx.max_row
    for c in range(1, 9):
        cell = ws_idx.cell(row=r, column=c)
        cell.font = bold_font
        cell.border = thin_border

if 'Historical 66 Capacity' not in existing_idx:
    ws_idx.append([
        'Historical 66 Capacity', 'STAFF_CAPACITY', 66,
        '29 Active Operational Accounts + 37 Unprovisioned Design Profiles',
        'Repository Planning Artifact', 'N/A (37 Profiles Have No Auth User)',
        '2026-10-21', 'Planning Baseline'
    ])
    r = ws_idx.max_row
    for c in range(1, 9):
        cell = ws_idx.cell(row=r, column=c)
        cell.font = bold_font
        cell.border = thin_border

wb.save('Documents/MediMind_Role_Credentials.xlsx')
print("Successfully updated Documents/MediMind_Role_Credentials.xlsx with Demonstration Accounts and 66 Capacity Roster.")
