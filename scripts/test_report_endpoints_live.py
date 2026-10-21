import requests
import json
import openpyxl
import io

# 1. Login as CHAIRMAN
wb = openpyxl.load_workbook('Documents/MediMind_Role_Credentials.xlsx', read_only=True)
ws = wb['Chairman']
headers = list(ws.iter_rows(values_only=True))[0]
row2 = list(ws.iter_rows(values_only=True))[1]
chair_email = row2[headers.index('Email / Username')]
chair_pw = row2[headers.index('Password')]

chair_login = requests.post(
    'http://localhost:5001/api/auth/login',
    json={'email': chair_email, 'password': chair_pw}
)
print(f"Chairman Login ({chair_email}): Status {chair_login.status_code}")
chair_token = chair_login.json()['data']['token']

headers = {
    'Authorization': f'Bearer {chair_token}'
}

# 2. Test GET /api/appointments/reports/preview via Gateway (Port 5000)
print("\n--- Testing Preview Endpoint via Gateway ---")
prev_res = requests.get('http://localhost:5000/api/appointments/reports/preview', headers=headers)
print(f"Preview Status: {prev_res.status_code}")
prev_data = prev_res.json()
print("Success:", prev_data.get('success'))
summary = prev_data.get('data', {}).get('summary', {})
print("Summary:", summary)
print("Total doctors in preview:", len(prev_data.get('data', {}).get('doctors', [])))
print("Total appointments returned in page 1:", len(prev_data.get('data', {}).get('appointments', [])))
print("Sample doctor 0:", prev_data.get('data', {}).get('doctors', [])[0])
print("Sample appointment 0:", prev_data.get('data', {}).get('appointments', [])[0])

# 3. Test GET /api/appointments/reports/export-excel via Gateway (Port 5000)
print("\n--- Testing Excel Export Endpoint via Gateway ---")
excel_res = requests.get('http://localhost:5000/api/appointments/reports/export-excel', headers=headers)
print(f"Excel Export Status: {excel_res.status_code}")
print(f"Content-Type: {excel_res.headers.get('Content-Type')}")
print(f"Content-Disposition: {excel_res.headers.get('Content-Disposition')}")
print(f"Excel Payload Size: {len(excel_res.content)} bytes")

# 4. Reopen and validate workbook with openpyxl
wb = openpyxl.load_workbook(io.BytesIO(excel_res.content))
print(f"\nReopened Exported Workbook! Sheets: {wb.sheetnames}")
for name in wb.sheetnames:
    ws = wb[name]
    print(f"  Sheet '{name}': rows={ws.max_row}, cols={ws.max_column}")

# Print first 3 rows of each sheet
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n--- Preview of '{name}' (first 3 rows) ---")
    for r in list(ws.iter_rows(values_only=True))[:3]:
        print(" ", r)

# 5. Security & Isolation test: Test that DOCTOR role CANNOT access reports (must return 403)
print("\n--- Testing Security / Role-Scoping (Doctor Role Prohibited) ---")
doc_login = requests.post(
    'http://localhost:5001/api/auth/login',
    json={'email': 'dr.rahul.sharma@aarogyam.hospital', 'password': 'Doctor#2026$01!Dev'}
)
doc_token = doc_login.json()['data']['token']
doc_headers = {'Authorization': f'Bearer {doc_token}'}

doc_prev = requests.get('http://localhost:5000/api/appointments/reports/preview', headers=doc_headers)
print(f"Doctor Preview Request Status: {doc_prev.status_code} (Must be 403 Forbidden)")
doc_export = requests.get('http://localhost:5000/api/appointments/reports/export-excel', headers=doc_headers)
print(f"Doctor Export Request Status: {doc_export.status_code} (Must be 403 Forbidden)")
